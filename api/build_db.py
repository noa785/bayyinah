"""
يبني قاعدة الأحاديث من ملفات Word اللي جمعتها هيفاء.

التشغيل:
    python build_db.py data/raw/*.docx

يطلع:
    data/hadiths.json   القاعدة اللي يقرأها النظام
    data/review.txt     الصفوف اللي تحتاج مراجعة هيفاء
"""

import json
import re
import sys
from pathlib import Path

from docx import Document

from arabic import normalize

# ---------- تصنيف الحكم (للون والقرار فقط، والعرض بلفظ الحكم كما ورد) ----------

DETAILED = ["المعنى", "الزياده", "الزيادة", "بهذا اللفظ", "رفعه"]
NOT_AUTHENTIC = ["موضوع", "باطل", "لا اصل", "لا أصل", "ليس له اصل", "ليس له أصل", "لا يصح",
                 "ليس بثابت", "ليس هو بثابت", "ليس بحديث", "منكر", "لا علم له", "ليس بصحيح", "لم يرد"]
WEAK = ["ضعيف", "مرسل"]
AUTHENTIC = ["صحيح", "حسن", "أخرجه في صحيحه", "اخرجه في صحيحه"]


def grade_class(grade: str) -> str:
    g = grade or ""
    if any(k in g for k in DETAILED):
        return "detailed"
    if any(k in g for k in NOT_AUTHENTIC):
        return "not_authentic"
    if any(k in g for k in WEAK):
        return "weak"
    if any(k in g for k in AUTHENTIC):
        return "authentic"
    return "unknown"


# ---------- تنظيف الخانات ----------

LABELS = re.compile(r"(المحدث|المصدر|الراوي|التخريج|خلاصة حكم المحدث)\s*:\s*")


def clean(s: str) -> str:
    s = (s or "").replace("\u200f", "").replace("\u200e", "").replace("\xa0", " ")
    s = re.sub(r"^[\s\-/|:]+", "", s)
    s = re.sub(r"\s*\n\s*", " ", s)
    s = re.sub(r"\s{2,}", " ", s)
    return s.strip(" /|-")


def parse_source(raw: str):
    """يفصل اسم الكتاب عن رقم الحديث او الجزء والصفحة."""
    s = LABELS.sub("", raw or "")
    parts = re.split(r"(?:ا?\s*لصفحة\s*[أا]و\s*الرقم|الرقم\s*[أا]و\s*ا?\s*لصفحة)\s*:?", s)
    book = clean(parts[0]) if parts else ""
    ref = clean(parts[1]) if len(parts) > 1 else ""
    ref = re.sub(r"\s*(التخريج|ال)\s*$", "", ref)
    return {"book": book or None, "ref": ref or None}


def field(raw: str, label: str):
    m = re.search(r"(?<!حكم )" + label + r"\s*:\s*([^|\n]+)", raw or "")
    return (clean(m.group(1)) or None) if m else None


def parse_alternative(raw: str):
    """البديل الصحيح مكتوب كنص ثم خلاصة الحكم والمحدث والمصدر."""
    raw = (raw or "").strip()
    if not raw:
        return None
    text = re.split(r"خلاصة حكم المحدث", raw)[0]
    grade = field(raw, "خلاصة حكم المحدث")
    if grade:
        grade = grade.strip("[] ")
    src_raw = raw.split("المصدر", 1)[1] if "المصدر" in raw else ""
    return {
        "text": clean(text),
        "grade_raw": grade,
        "graded_by": field(raw, "المحدث"),
        "narrator": field(raw, "الراوي") or None,
        "source": parse_source(src_raw) if src_raw else {"book": None, "ref": None},
        "url": None,
    }


# ---------- قراءة الجداول ----------

COLS = {
    "text": ["النص المنتشر", "نص الحديث"],
    "grade": ["نوعه", "حكمه"],
    "graded_by": ["من حكم عليه", "المحدث"],
    "source": ["المصدر"],
    "alternative": ["البديل الصحيح"],
    "takhrij": ["التخريج", "االتخريج"],
    "narrator": ["الراوي"],
}


def map_header(cells):
    m = {}
    for i, c in enumerate(cells):
        c = c.strip()
        for key, names in COLS.items():
            if c in names and key not in m:
                m[key] = i
    return m if "text" in m and "grade" in m else None


def strip_leading_empty(cells):
    """بعض الصفوف فيها خانة فاضية زايدة في البداية، فنشيلها عشان تتطابق مع العناوين."""
    while len(cells) > 1 and not cells[0].strip() and any(x.strip() for x in cells[1:]):
        cells = cells[1:]
    return cells


def read_docx(path: Path, review: list):
    doc = Document(path)
    rows, header = [], None
    for t in doc.tables:
        for r in t.rows:
            raw = [c.text for c in r.cells]
            cells = strip_leading_empty(raw)
            h = map_header(cells)
            if h:
                header = h
                continue
            if header is None or not any(x.strip() for x in cells):
                continue
            shifted = len(cells) != len(raw) and len(raw) > max(header.values()) + 1
            get = lambda k: cells[header[k]] if k in header and header[k] < len(cells) else ""
            rec = {
                "text": clean(get("text")),
                "grade_raw": clean(LABELS.sub("", get("grade"))).strip("[]"),
                "graded_by": clean(LABELS.sub("", get("graded_by"))),
                "source": parse_source(get("source")),
                "takhrij": clean(LABELS.sub("", get("takhrij"))) or None,
                "narrator": clean(LABELS.sub("", get("narrator"))) or None,
                "alternative": parse_alternative(get("alternative")),
                "file": path.name,
            }
            if not rec["text"]:
                continue
            if shifted:
                review.append(f"[{path.name}] كانت خاناته مزاحة وتعدلت تلقائيا، تأكدي منه: {rec['text'][:60]}")
            rows.append(rec)
    return rows


def tidy_name(s):
    """يصلح أخطاء الكتابة في أسماء الكتب والعلماء الجاية من ملفات الوورد."""
    if not s:
        return s
    s = re.sub(r"\s+", " ", s).strip(" []،,")
    s = re.sub(r"^اال", "ال", s)
    s = re.sub(r"\s+(ال|ا)$", "", s)
    words = s.split()
    half = len(words) // 2
    if len(words) % 2 == 0 and half and words[:half] == words[half:]:
        s = " ".join(words[:half])
    return {"الالباني": "الألباني"}.get(s, s)


def tidy(rec):
    if not rec:
        return
    for k in ("graded_by", "narrator"):
        rec[k] = tidy_name(rec.get(k))
    if rec.get("source") and rec["source"].get("book"):
        rec["source"]["book"] = tidy_name(rec["source"]["book"])


ALT_RE = re.compile(r"^(.*?)\s*\(\s*([^،()]+)،\s*(.+?)\s+([\d/٠-٩]+)\s*\)\s*$")


def read_xlsx(path: Path, review: list):
    """يقرأ جدول المراجعة الشرعية، ويدخل فقط الصفوف اللي اعتمدتها المتخصصة."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        head = [str(c.value or "").strip() for c in ws[1]]
        if "النص المنتشر" not in head or "قرارك" not in head:
            continue
        col = {h: i for i, h in enumerate(head)}
        for r in ws.iter_rows(min_row=2, values_only=True):
            g = lambda h: clean(str(r[col[h]] or "")).strip() if h in col else ""
            if g("قرارك") != "اعتماد" or not g("النص المنتشر"):
                continue
            alt, raw = None, g("البديل الصحيح")
            if raw:
                m = ALT_RE.match(raw)
                if m:
                    alt = {"text": m.group(1), "grade_raw": m.group(2), "graded_by": None, "narrator": None,
                           "source": {"book": m.group(3), "ref": m.group(4)}, "url": None}
                else:
                    review.append(f"[{path.name}] البديل ما انقرأ، تأكدي من صيغته: {raw[:60]}")
            rows.append({
                "text": g("النص المنتشر"), "grade_raw": g("نوعه"), "graded_by": g("من حكم عليه") or None,
                "source": {"book": g("المصدر") or None, "ref": g("الرقم أو الصفحة") or None},
                "takhrij": None, "narrator": None, "alternative": alt, "note": None,
                "file": path.name, "reviewed": True,
            })
    return rows


def apply_corrections(rows, review, path="data/corrections.json"):
    """يطبق مراجعات المتخصصة الشرعية على البيانات، والملفات الأصلية تبقى كما هي."""
    f = Path(path)
    if not f.exists():
        return rows
    fixes = json.loads(f.read_text(encoding="utf-8"))
    removes = [normalize(r["starts_with"]) for r in fixes.get("remove", [])]
    out = []
    for rec in rows:
        key = normalize(rec["text"])
        if any(key.startswith(r) for r in removes):
            review.append(f"انحذف بقرار المراجعة الشرعية: {rec['text'][:60]}")
            continue
        for u in fixes.get("update", []):
            if key.startswith(normalize(u["starts_with"])):
                rec.update(u["set"])
                rec["reviewed"] = True
        out.append(rec)
    # أحاديث جديدة أضافتها المتخصصة الشرعية بعد المراجعة
    for a in fixes.get("add", []):
        rec = {"takhrij": None, "narrator": None, "alternative": None, **a, "file": "corrections.json", "reviewed": True}
        out.append(rec)
    return out


def main(paths):
    review, all_rows = [], []
    for p in paths:
        all_rows.extend(read_xlsx(Path(p), review) if str(p).endswith(".xlsx") else read_docx(Path(p), review))
    all_rows = apply_corrections(all_rows, review)
    for rec in all_rows:
        tidy(rec)
        tidy(rec.get("alternative"))

    db, seen = [], set()
    for rec in all_rows:
        key = normalize(rec["text"])
        if key in seen:
            review.append(f"مكرر وانحذف: {rec['text'][:60]}")
            continue
        seen.add(key)
        rec["id"] = f"H{len(db) + 1:03d}"
        rec["grade_class"] = grade_class(rec["grade_raw"])
        rec["note"] = rec.get("note")
        rec["url"] = None
        if rec["grade_class"] in ("unknown", "detailed") and not rec.get("reviewed"):
            review.append(f"{rec['id']} الحكم يحتاج تأكيد ({rec['grade_raw']}): {rec['text'][:60]}")
        db.append(rec)

        # البديل الصحيح ندخله كحديث مستقل عشان لو أحد قاله يطلع مؤيد
        alt = rec["alternative"]
        if alt and alt["text"]:
            akey = normalize(alt["text"])
            if akey not in seen:
                seen.add(akey)
                db.append({
                    "id": rec["id"] + "-ALT",
                    "text": alt["text"],
                    "grade_raw": alt["grade_raw"] or "",
                    "grade_class": grade_class(alt["grade_raw"] or ""),
                    "graded_by": alt["graded_by"],
                    "source": alt["source"],
                    "takhrij": None,
                    "narrator": alt["narrator"],
                    "alternative": None,
                    "note": None,
                    "url": None,
                    "file": rec["file"],
                })

    out = Path("data")
    out.mkdir(exist_ok=True)
    (out / "hadiths.json").write_text(json.dumps(db, ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "review.txt").write_text("\n".join(review), encoding="utf-8")

    counts = {}
    for r in db:
        counts[r["grade_class"]] = counts.get(r["grade_class"], 0) + 1
    print(f"تم: {len(db)} حديث في data/hadiths.json")
    print("التوزيع:", counts)
    print(f"ملاحظات للمراجعة: {len(review)} في data/review.txt")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("الاستخدام: python build_db.py data/raw/*.docx")
        sys.exit(1)
    main(sys.argv[1:])
